"""[EXPORT] Real XLSX, scoped on the server, with text cells never interpreted as formulas."""
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
import mysql_store
from roles import scoped

SHEETS = {
 'inspections': ('점검 기록', [('inspected_at','점검 일시'),('equipment_id','설비 ID'),('inspector','점검자'),('model_status','판정 상태'),('operation_state','운전 상태'),('noise','소음'),('vibration','진동'),('heat','발열'),('decision','대응 결정'),('memo','메모'),('approval_status','일정 승인 상태')]),
 'work_history': ('정비 작업', [('work_date','작업일'),('equipment_id','설비 ID'),('work_type','작업 구분'),('anomaly','이상 내용'),('action','조치'),('result','결과'),('worker','작업자')]),
 'maintenance_reviews': ('정비 검토', [('registered_at','등록일'),('equipment_id','설비 ID'),('target','대상'),('content','내용'),('priority','우선순위'),('review_status','검토 상태'),('approval_status','일정 승인 상태')]),
 'follow_ups': ('점검 일정', [('due_date','예정일'),('equipment_id','설비 ID'),('type','종류'),('assignee','담당자'),('items','내용'),('status','상태')]),
}


def workbook(user, process='all', equipment='all'):
    data = scoped(mysql_store.read_tables(list(SHEETS)), user)
    eq = scoped(mysql_store.read_tables(['equipment']), user)['equipment']
    ids = {r['equipment_id'] for r in eq if process in ('all', r['process']) and equipment in ('all', r['equipment_id'])}
    names = {r['equipment_id']: r['equipment_name'] for r in eq}
    wb = Workbook(); wb.remove(wb.active)
    for key, (title, columns) in SHEETS.items():
        ws = wb.create_sheet(title)
        ws.append([label for _, label in columns] + ['설비명'])
        for row in data[key]:
            if row['equipment_id'] in ids:
                def value(k):
                    raw = row.get(k, '')
                    return {'pending':'승인 대기','approved':'승인','rejected':'반려'}.get(raw, '기존 기록') if k == 'approval_status' else raw
                ws.append([value(k) for k, _ in columns] + [names[row['equipment_id']]])
                for cell in ws[ws.max_row]:
                    if isinstance(cell.value, str):
                        cell.data_type = 's'  # Includes strings beginning with =, +, -, @.
                    cell.alignment = Alignment(vertical='top', wrap_text=True)
        for cell in ws[1]:
            cell.font = Font(color='FFFFFF', bold=True)
            cell.fill = PatternFill('solid', fgColor='00547A')
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = 26
        ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
    output = BytesIO(); wb.save(output)
    return output.getvalue()
